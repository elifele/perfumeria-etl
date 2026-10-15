import pandas as pd
from datetime import date
import sys 
import os
from dotenv import load_dotenv
from sqlalchemy import create_engine
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.append(r"C:\Users\Admin\OneDrive\Documentos\mis_funciones")
from etl_funciones import(
    cargar_excel,
    validar_calidad,
    limpiar_dataframe,
    cargar_postgresql,
    ejecutar_query,
    exportar_excel,
    enviar_email
)

#CONECTAR A POSTGRESQL
load_dotenv()

cargar_postgresql(
    database=os.getenv("DB_NAME"),
    user=os.getenv("DB_USER"),
    password=os.getenv("DB_PASSWORD"),
)
db_url = (
    f"postgresql://{os.getenv('DB_USER')}:"
    f"{os.getenv('DB_PASSWORD')}@"
    f"{os.getenv('DB_HOST')}:"
    f"{os.getenv('DB_PORT')}/"
    f"{os.getenv('DB_NAME')}"
)

engine = create_engine(db_url)

#CARGAR EXCEL
df = cargar_excel("venta_perfumeria.xlsx")

print(df.head())
print(df.columns.to_list())

#VALIDAR CALIDAD
validar_general = validar_calidad(df)
validar_precio = validar_calidad(df, "Precio_unitario")

print("\nNULOS")
print(validar_general["nulos"])

print("\nDUPLICADOS")
print(validar_general["duplicados"])

print("\nVALORES INAVALIDOS (Precio_unitario)")
print(validar_precio["valores_invalidos"])

#LIMPIAR DATAFRAME
df = limpiar_dataframe(
    df,
    columnas_texto=["Nombre_producto","Categoria", "Canal_venta", "Metodo_pago"],
    columnas_fecha=["Fecha"]
)
print(df.head())
print("limpieza completada")

#CARGAR A POSTGRESQL
df.to_sql("ventas_perfumeria", engine, index=False, if_exists="replace")

print("tabla cargada a postgresql")

#EJECUTAR QUERY 
#QUERY 1 Ventas por categoria
df_categoria = pd.read_sql("""
    SELECT "Categoria",
        COUNT(*) AS "Cantidad_ventas",
        SUM("Cantidad") AS "Unidades_vendidas",
        ROUND(SUM("Monto_total")::numeric, 2) AS "Ingresos_totales"
    FROM ventas_perfumeria
    GROUP BY "Categoria"
    ORDER BY "Ingresos_totales" DESC
""",engine)

print("\nventas_categoria")
print(df_categoria)

#QUERY 2 Top 5 productos
df_productos = pd.read_sql("""
    SELECT "Nombre_producto", "Categoria",
        SUM("Cantidad") AS "Unidades_vendidas",
        ROUND(SUM("Monto_total")::numeric, 2) AS "Ingresos_totales"
    FROM ventas_perfumeria
    GROUP BY "Nombre_producto" , "Categoria"
    ORDER BY "Unidades_vendidas" DESC
    LIMIT 5
""",engine)

print("\ntop_5_productos")
print(df_productos) 

#QUERY 3 Ventas por canal
df_canal_venta = pd.read_sql("""
    SELECT "Canal_venta",
        COUNT(*) AS "Cantidad_ventas",
        ROUND(SUM("Monto_total")::numeric, 2) AS "ingresos_totales"
    FROM ventas_perfumeria
    GROUP BY "Canal_venta"
    ORDER BY "ingresos_totales" DESC
""",engine)

print("\nventas_canal")
print(df_canal_venta)

#QUERY 4 Metodo de pago mas usado 
df_metodo_pago = pd.read_sql("""
    SELECT "Metodo_pago",
        COUNT(*) AS "Cantidad_ventas",
        ROUND(SUM("Monto_total")::numeric, 2) AS "Ingresos_totales"
    FROM ventas_perfumeria
    GROUP BY "Metodo_pago"
    ORDER BY "Ingresos_totales" DESC
""",engine)

print("\nmetodo_pago")
print(df_metodo_pago)

#EXPORTAR A EXCEL
df["Fecha"] = pd.to_datetime(df["Fecha"]).dt.date
hoy = date.today().strftime("%Y%m%d")

exportar_excel(f"reporte_venta_perfumeria_{hoy}.xlsx",{
    "datos_completos" : df,
    "ventas_categoria": df_categoria,
    "top_5_productos" : df_productos,
    "ventas_canal"    : df_canal_venta,
    "metodo_pago"     : df_metodo_pago
})

#FORMATO PROFESIONAL 
wb = load_workbook(f"reporte_venta_perfumeria_{hoy}.xlsx")

azul_oscuro = "1F4E79"
azul_medio  = "2E75B6"
azul_claro  = "D6E4F0"
blanco      = "FFFFFF"

borde_fino = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC')
)

for nombre_hoja in wb.sheetnames:
    ws = wb[nombre_hoja]

    #Encabezados fila 1
    for cell in ws["1"]:
        cell.font      = Font(name="Arial", size=9, bold=True, color=blanco)
        cell.fill      = PatternFill("solid", fgColor=azul_oscuro)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = borde_fino
    ws.row_dimensions[1].height = 35

    #Filas datos
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        color = azul_claro if i % 2 == 0 else blanco
        for j, cell in enumerate(row, start=1):
            cell.fill   = PatternFill("solid", fgColor=color)
            cell.border = borde_fino
            cell.font   = Font(name="Arial", size=9)
            if j == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = Font(name="Arial", size=9, bold=True)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, int):
                    cell.number_format = '#,##0'
                
                elif isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
        ws.row_dimensions[i].height = 18

    #Ancho automatico 
    for col in ws.iter_cols():
        max_len = 0
        col_letter = None
        for cell in col:
            if hasattr(cell, 'column_letter'):
                col_letter = cell.column_letter
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        if col_letter:
            ws.column_dimensions[col_letter].width = min(max_len + 2, 15)

     # Ajustar ancho según encabezado
    for cell in ws[1]:
        if cell.value:
            col_letter = cell.column_letter
            ancho_encabezado = len(str(cell.value)) + 2
            ancho_actual = ws.column_dimensions[col_letter].width
            ws.column_dimensions[col_letter].width = max(ancho_actual, ancho_encabezado)

wb.save(f"reporte_venta_perfumeria_{hoy}.xlsx")
print("Formato aplicado correctamente")

#ENVIAR EMAIL
enviar_email(
    destinatario=os.getenv("EMAIL_USER"),
    asunto="Reporte de ventas perfumeria",
    mensaje="Adjunto el reporte de ventas perfumeria actualizado",
    archivo=f"reporte_venta_perfumeria_{hoy}.xlsx",
    user=os.getenv("EMAIL_USER"),
    password=os.getenv("EMAIL_PASSWORD")
)









